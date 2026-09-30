#!/usr/bin/env python3
"""extract_ufgs_text.py — estrae il testo dalle specifiche UFGS (pdf e zip/docx).
Output: raw/ufgs_specs/txt/<nome>.txt per ogni specifica."""
import glob, os, sys, zipfile, io, time

HERE = os.path.dirname(os.path.abspath(__file__))
RAW = os.path.join(HERE, "raw", "ufgs_specs")
OUT = os.path.join(RAW, "txt")
os.makedirs(OUT, exist_ok=True)

def pdf_text(path):
    from pypdf import PdfReader
    r = PdfReader(path)
    return "\n\n".join((p.extract_text() or "") for p in r.pages)

def docx_text(blob):
    import docx
    d = docx.Document(io.BytesIO(blob))
    parts = [p.text for p in d.paragraphs]
    for tb in d.tables:
        for row in tb.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)

def doc_text_soffice(path):  # .doc legacy: non gestito, segnalato
    return None

files = sorted(glob.glob(os.path.join(RAW, "*.pdf")) + glob.glob(os.path.join(RAW, "*.zip")))
ok, fail = 0, []
for f in files:
    base = os.path.splitext(os.path.basename(f))[0]
    dest = os.path.join(OUT, base + ".txt")
    if os.path.exists(dest) and os.path.getsize(dest) > 500:
        ok += 1; continue
    try:
        if f.lower().endswith(".pdf"):
            t = pdf_text(f)
        else:
            z = zipfile.ZipFile(f)
            parts = []
            for n in z.namelist():
                if n.lower().endswith(".docx"):
                    parts.append(f"### FILE INTERNO: {n}\n" + docx_text(z.read(n)))
                elif n.lower().endswith(".pdf"):
                    tmp = os.path.join(OUT, "_tmp.pdf")
                    open(tmp, "wb").write(z.read(n))
                    parts.append(f"### FILE INTERNO: {n}\n" + pdf_text(tmp))
                    os.remove(tmp)
                elif n.lower().endswith(".doc"):
                    parts.append(f"### FILE INTERNO (formato .doc legacy, non estratto): {n}")
            t = "\n\n".join(parts)
        open(dest, "w", encoding="utf-8").write(t)
        ok += 1
    except Exception as e:
        fail.append((os.path.basename(f), str(e)[:100]))

print(f"estratte: {ok}/{len(files)}")
for n, e in fail: print(f"  [!!] {n}: {e}")
tot = sum(os.path.getsize(p) for p in glob.glob(os.path.join(OUT, "*.txt")))
print(f"testo totale: {tot:,} byte")
